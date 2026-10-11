import csv
import io
import json
import tempfile
import zipfile
from copy import deepcopy
from datetime import date, time
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
import httpx
from openpyxl import load_workbook
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError, PermissionDenied
from django.db import transaction, connection, DatabaseError
from django.test import TestCase, SimpleTestCase, override_settings, Client
from django.utils import timezone
from apps.collisions.models import Collision
from apps.collisions.services import save_entity
from apps.core.models import SavedQuery, DataRevision
from apps.analytics import risk as risk_services
from apps.analytics.forms import make_form as analysis_form
from apps.analytics.models import RiskRun
from apps.forecasting import services as forecast_services, jobs as forecast_jobs
from apps.forecasting.forms import make_form as forecast_form
from apps.reports import services, exports, ai, jobs
from apps.reports.forms import prepare_forms
from apps.reports.models import Report, ReportSnapshot, ExportJob, AIBrief, AIRequest, BriefRevision
from apps.reports.tables import memory_table, tables
from apps.reports.figures import svg
from tests import test_s06 as forecast_fixture
from tests import test_s04 as flow_fixture

CONFIG=ai.Config('deepseek','deepseek-flash','https://api.deepseek.com/chat/completions','unit-only-secret')
HTTP_CLIENT=httpx.Client


def valid_output(summary):
    ids=[r['id'] for r in summary['facts'][:2]]
    return json.dumps({'title':'交通安全分析简报','findings':[{'text':'所选范围的登记量应按各自统计单位理解。','refs':ids}],
        'interpretations':[{'text':'结论受来源覆盖及已报告字段的完整性限制。','refs':['L001']}],
        'recommendations':[{'text':'建议核查相应来源与缺失记录，不能据此认定原因。','refs':ids[:1]+['L001']}]},ensure_ascii=False)


class ExportAndAIUnitTests(SimpleTestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='traffic-s07-')
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.summary={'facts':[{'id':'F001','text':'事故：0 起'},{'id':'F002','text':'已报告受伤：2 人'}],
                      'limitations':[{'id':'L001','text':'来源截止有明确边界。'}]}

    def test_csv_splitting_preserves_all_rows_zero_and_formula_text(self):
        rows=[{'x':x,'y':'=1+1' if x==0 else None} for x in range(5)]
        table=memory_table('rows','全部记录',[('x','数值'),('y','文字')],rows)
        output=self.root/'data.zip'
        with patch.object(exports,'CSV_DATA_ROWS',2):
            counts=exports.write_csv(output,[table],{'notes':['测试范围']})
        self.assertEqual(counts,{'全部记录':5})
        with zipfile.ZipFile(output) as z:
            names=[n for n in z.namelist() if n!='00_说明.csv']
            self.assertEqual(len(names),3)
            data=[r for name in names for r in list(csv.reader(io.StringIO(z.read(name).decode('utf-8-sig'))))[1:]]
        self.assertEqual([r[0] for r in data],['0','1','2','3','4'])
        self.assertEqual(data[0][1],"'=1+1");self.assertEqual(data[1][1],'')

    def test_xlsx_splits_without_loss_and_formulas_are_string_cells(self):
        rows=[{'v':i,'s':'=HYPERLINK("x")' if i==0 else None} for i in range(5)]
        path=self.root/'data.xlsx'
        with patch.object(exports,'XLSX_DATA_ROWS',2):
            exports.write_xlsx(path,[memory_table('rows','记录',[('v','数值'),('s','文字')],rows)],{'notes':[]})
        wb=load_workbook(path,read_only=True,data_only=False)
        try:
            self.assertEqual(len(wb.sheetnames),4)
            data=[r for sheet in wb.worksheets[1:] for r in sheet.iter_rows(min_row=2)]
            self.assertEqual([r[0].value for r in data],[0,1,2,3,4])
            self.assertEqual(data[0][1].data_type,'s');self.assertEqual(data[0][1].value,'=HYPERLINK("x")')
            self.assertIsNone(data[1][1].value)
        finally:wb.close()

    def test_empty_csv_retains_headers_and_scope_metadata(self):
        path=self.root/'empty.zip'
        exports.write_csv(path,[memory_table('empty','空范围',[('v','数值')],[])],{'conditions':{'borough':'5'},'notes':[]})
        with zipfile.ZipFile(path) as z:
            self.assertEqual(list(csv.reader(io.StringIO(z.read('01_empty_001.csv').decode('utf-8-sig')))),[['数值']])
            self.assertIn('"borough"',z.read('00_说明.csv').decode('utf-8-sig'))

    def test_excel_long_text_is_rejected_without_silent_truncation(self):
        from openpyxl import Workbook
        sheet=Workbook(write_only=True).create_sheet('test')
        with self.assertRaisesRegex(exports.ExportError,'excel_cell_too_long'):exports.excel_cell(sheet,'字'*32768)
        with self.assertRaisesRegex(exports.ExportError,'invalid_text'):exports.excel_cell(sheet,'x\x00y')

    def test_formula_leading_whitespace_is_escaped_but_numeric_negative_is_not(self):
        for value in ['=x',' +x','\t-x','@x','\n=x']:
            self.assertTrue(exports.csv_value(value).startswith("'"))
        self.assertEqual(exports.csv_value(-2),'-2')

    def test_svg_escapes_labels_and_preserves_missing_gaps(self):
        chart={'kind':'line','title':'<script>bad</script>','unit':'起','labels':['a','b','c'],'series':[{'name':'数量','values':[0,None,5]}],'note':''}
        data=svg(chart)
        self.assertNotIn('<script>',data);self.assertIn('&lt;script&gt;',data)
        self.assertEqual(data.count('<polyline'),2)

    def test_config_rejects_credential_destinations_and_does_not_echo_key(self):
        path=self.root/'private.json'
        base={'provider':'deepseek','model':'deepseek-flash','api_key':'unit-only-secret'}
        for url in ['http://api.deepseek.com','https://api.deepseek.com.evil.test','https://x@api.deepseek.com','https://api.deepseek.com?key=x']:
            path.write_text(json.dumps({**base,'base_url':url}),encoding='utf-8')
            with override_settings(AI_CONFIG_FILE=str(path)),self.assertRaises(ai.AIError) as raised:ai.load_config()
            self.assertNotIn(base['api_key'],str(raised.exception))
        path.write_text(json.dumps({**base,'base_url':'https://api.deepseek.com'}),encoding='utf-8')
        with override_settings(AI_CONFIG_FILE=str(path)):self.assertNotIn(base['api_key'],repr(ai.load_config()))

    def test_valid_structured_output_uses_only_authoritative_numeric_references(self):
        text=ai.parse_output(valid_output(self.summary),self.summary)
        self.assertIn('事故：0 起',text);self.assertIn('已报告受伤：2 人',text)
        self.assertIn('依据：',text)

    def test_unsupported_numbers_references_markup_and_empty_content_are_rejected(self):
        good=json.loads(valid_output(self.summary))
        variants=[]
        for content in ['新增事故 999 起','<script>alert(x)</script>','http://unknown.test']:
            bad=deepcopy(good);bad['findings'][0]['text']=content;variants.append(json.dumps(bad))
        bad=deepcopy(good);bad['findings'][0]['refs']=['F999'];variants.append(json.dumps(bad))
        variants+=['', '[]', '{"title":"不完整"}']
        for raw in variants:
            with self.assertRaises(ai.AIError):ai.parse_output(raw,self.summary)

    def test_http_errors_do_not_expose_response_or_follow_redirects(self):
        for status,code in [(401,'authentication'),(402,'quota'),(429,'rate_limit'),(503,'provider'),(307,'provider')]:
            seen=[]
            def handler(request):
                seen.append(request)
                return httpx.Response(status,headers={'Location':'https://untrusted.test'},text='unit-only-secret')
            with patch('apps.reports.ai.httpx.Client',side_effect=lambda **kw:HTTP_CLIENT(transport=httpx.MockTransport(handler),**kw)),self.assertRaises(ai.AIError) as raised:
                ai.call_model(CONFIG,self.summary)
            self.assertEqual(raised.exception.code,code);self.assertEqual(len(seen),1)
            self.assertNotIn('unit-only-secret',str(raised.exception))

    def test_http_timeout_and_truncated_json_are_distinct_failures(self):
        def timeout(request):raise httpx.ReadTimeout('untrusted sensitive message')
        with patch('apps.reports.ai.httpx.Client',side_effect=lambda **kw:HTTP_CLIENT(transport=httpx.MockTransport(timeout),**kw)),self.assertRaises(ai.AIError) as e:ai.call_model(CONFIG,self.summary)
        self.assertEqual(e.exception.code,'timeout')
        def partial(request):return httpx.Response(200,json={'choices':[{'finish_reason':'length','message':{'content':'{}'}}]})
        with patch('apps.reports.ai.httpx.Client',side_effect=lambda **kw:HTTP_CLIENT(transport=httpx.MockTransport(partial),**kw)),self.assertRaises(ai.AIError) as e:ai.call_model(CONFIG,self.summary)
        self.assertEqual(e.exception.code,'format')


@override_settings(ALLOWED_HOSTS=['testserver','127.0.0.1','localhost'],AI_CONFIG_FILE='')
class ReportDatabaseTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        forecast_fixture.ForecastDatabaseTests.setUpTestData.__func__(cls)

    def setUp(self):
        self.client.force_login(self.analyst)
        self.tmp=tempfile.TemporaryDirectory(prefix='traffic-s07-db-')
        self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)

    def report(self,kind='statistics',**values):
        main,source=prepare_forms(self.analyst,{'kind':kind,'date_from':'2021-01-01','date_to':'2025-12-31','network':str(self.network.pk),**values})
        self.assertTrue(main.is_valid(),main.errors)
        if source:self.assertTrue(source.is_valid(),source.errors)
        report=services.enqueue(self.analyst,main,source)
        self.assertEqual(jobs.claim('report'),report.pk)
        self.assertTrue(jobs.execute('report',report.pk))
        report.refresh_from_db();self.assertEqual(report.status,'succeeded')
        return report

    def export(self,fmt='csv',report=None,**conditions):
        form=exports.make_form('crashes',conditions) if not report else None
        job=exports.enqueue(self.analyst,'report' if report else 'crashes',fmt,form=form,report=report)
        self.assertEqual(jobs.claim('export'),job.pk)
        with override_settings(BASE_DIR=self.path):
            self.assertTrue(jobs.execute('export',job.pk))
        job.refresh_from_db();self.assertEqual(job.status,'succeeded')
        return job

    def brief(self,report=None):
        report=report or self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            brief,_=ai.enqueue(self.analyst,report)
            self.assertEqual(jobs.claim('brief'),brief.pk)
            raw=valid_output(brief.input_summary)
            with patch('apps.reports.ai.call_model',return_value=(raw,ai.parse_output(raw,brief.input_summary),{'prompt_tokens':50,'completion_tokens':20},10)):
                jobs.execute('brief',brief.pk)
        brief.refresh_from_db();self.assertEqual(brief.status,'succeeded')
        return brief

    def test_stat_report_freezes_actual_counts_and_does_not_change_after_fact_edit(self):
        report=self.report()
        n=next(m['value'] for m in report.payload['metrics'] if m['key']=='n')
        self.assertEqual(n,len(self.events))
        snapshot=deepcopy(report.payload)
        event=Collision.objects.get(pk=self.events[0])
        save_entity(kind='crashes',actor=self.editor,pk=event.pk,expected_revision=event.revision,data={'persons_injured':9},reason='隔离测试验证历史报告不随当前事实修订。')
        report.refresh_from_db();self.assertEqual(report.payload,snapshot)
        self.assertEqual(ReportSnapshot.objects.filter(report=report).count(),2)

    def test_comparison_report_keeps_zero_baseline_and_period_limits(self):
        report=self.report('comparison',date_from='2021-01-01',date_to='2021-01-10',compare_from='2025-01-11',compare_to='2025-01-31')
        row=report.payload['analysis']['comparisons'][0]
        self.assertEqual(row['a'],0);self.assertIsNone(row['percent']);self.assertTrue(row['reasons'])

    def test_forecast_report_references_frozen_full_backtest(self):
        form=forecast_form({'date_from':'2021-01-01','date_to':'2025-12-31','network':str(self.network.pk)})
        run=forecast_services.enqueue(self.analyst,form);forecast_jobs.claim();forecast_jobs.execute(run.pk)
        report=self.report('forecast',forecast_run=str(run.pk))
        counts={t.code:t.count for t in tables(report)}
        self.assertEqual(counts['history'],60);self.assertEqual(counts['backtest'],run.folds.count());self.assertGreater(counts['backtest'],30)
        self.assertEqual(len(report.payload['charts'][0]['labels']),63)

    def test_risk_report_uses_all_immutable_results_not_preview(self):
        form=analysis_form('risk',{'date_from':'2021-01-01','date_to':'2025-12-31','level':'road','network':str(self.network.pk)})
        run=risk_services.enqueue(self.analyst,form);RiskRun.objects.filter(pk=run.pk).update(status='running');risk_services.compute_run(run.pk)
        report=self.report('risk',risk_run=str(run.pk))
        table=list(tables(report))[0]
        self.assertEqual(table.count,run.results.count())

    def test_report_queries_reuse_only_owned_saved_conditions(self):
        saved=SavedQuery.objects.create(owner=self.analyst,name='年度统计',kind='statistics',schema_version=2,conditions={'date_from':'2025-01-01','date_to':'2025-12-31','borough':'1'})
        main,source=prepare_forms(self.analyst,{'saved_query':str(saved.pk)},load_saved=True)
        self.assertTrue(main.is_valid());self.assertEqual(source.data['borough'],'1')
        with self.assertRaises(ValidationError):prepare_forms(self.other,{'saved_query':str(saved.pk)},load_saved=True)

    def test_completed_report_and_sources_are_database_immutable(self):
        report=self.report()
        for sql in ["UPDATE reports_report SET payload='{}' WHERE id=%s","DELETE FROM reports_report WHERE id=%s"]:
            with self.assertRaises(DatabaseError),transaction.atomic(),connection.cursor() as cursor:cursor.execute(sql,[report.pk])
        with self.assertRaises(DatabaseError),transaction.atomic():ReportSnapshot.objects.filter(report=report).delete()

    def test_report_failure_rolls_back_partial_sources_and_payload(self):
        main,source=prepare_forms(self.analyst,{'kind':'statistics'})
        report=services.enqueue(self.analyst,main,source);jobs.claim('report')
        with patch('apps.reports.services.ReportSnapshot.objects.bulk_create',side_effect=RuntimeError('fault')):self.assertFalse(jobs.execute('report',report.pk))
        report.refresh_from_db();self.assertEqual(report.status,'failed');self.assertEqual(report.payload,{});self.assertFalse(report.sources.exists())

    def test_full_csv_matches_filter_beyond_first_page(self):
        job=self.export(date_from='2025-01-01',date_to='2025-12-31',borough='1')
        expected=Collision.objects.filter(local_date__year=2025,borough_id=1).count()
        self.assertGreater(expected,30);self.assertEqual(job.row_count,expected)
        with override_settings(BASE_DIR=self.path):
            with zipfile.ZipFile(exports.path_for(job)) as z:
                rows=list(csv.DictReader(io.StringIO(z.read('01_crashes_001.csv').decode('utf-8-sig'))))
        self.assertEqual(len(rows),expected);self.assertEqual({r['纽约当地日期'][:4] for r in rows},{'2025'})
        self.assertEqual(job.metadata['conditions']['borough'],'1')

    def test_report_xlsx_has_summary_sources_and_every_table(self):
        report=self.report()
        job=self.export('xlsx',report=report)
        with override_settings(BASE_DIR=self.path):wb=load_workbook(exports.path_for(job),read_only=True)
        try:
            self.assertEqual(wb.sheetnames[0],'说明')
            self.assertEqual(sum(1 for _ in wb.worksheets[1].iter_rows())-1,60)
        finally:wb.close()
        self.assertEqual(job.row_count,60)

    def test_preexisting_export_artifact_is_never_deleted_on_failure(self):
        form=exports.make_form('crashes',{})
        job=exports.enqueue(self.analyst,'crashes','csv',form=form);jobs.claim('export')
        with override_settings(BASE_DIR=self.path):
            path=exports.path_for(job);path.parent.mkdir(parents=True);path.write_bytes(b'preserve-existing')
            self.assertFalse(jobs.execute('export',job.pk))
            self.assertEqual(path.read_bytes(),b'preserve-existing')

    def test_export_failure_does_not_publish_partial_file(self):
        form=exports.make_form('crashes',{})
        job=exports.enqueue(self.analyst,'crashes','csv',form=form);jobs.claim('export')
        with override_settings(BASE_DIR=self.path),patch('apps.reports.exports.write_csv',side_effect=RuntimeError('disk failure')):
            self.assertFalse(jobs.execute('export',job.pk));self.assertFalse(exports.path_for(job).exists())
            self.assertEqual(list((self.path/'var/exports').glob('*.part')),[])

    def test_permissions_and_owner_checks_cover_every_private_route(self):
        report=self.report();brief=self.brief(report);job=self.export(report=report)
        paths=[f'/reports/{report.pk}/',f'/exports/{job.pk}/',f'/exports/{job.pk}/download/',f'/briefs/{brief.pk}/',f'/briefs/{brief.pk}/edit/',f'/briefs/{brief.pk}/download/',f'/report-jobs/report/{report.pk}/']
        self.client.force_login(self.other)
        for path in paths:self.assertEqual(self.client.get(path).status_code,404,path)
        nobody=User.objects.create_user('report-no-role');self.client.force_login(nobody)
        self.assertEqual(self.client.get('/reports/').status_code,403)
        with self.assertRaises(PermissionDenied):services.allow(nobody)

    def test_normal_pages_render_when_model_configuration_is_absent(self):
        report=self.report();brief=self.brief(report)
        for path in ['/reports/','/reports/new/','/exports/','/exports/new/',f'/reports/{report.pk}/',f'/briefs/{brief.pk}/',f'/briefs/{brief.pk}/edit/']:
            self.assertEqual(self.client.get(path).status_code,200,path)
        self.assertEqual(self.client.get('/exports/new/?saved_query=not-a-number').status_code,404)
        self.assertEqual(self.client.get('/reports/new/?saved_query=999999').status_code,400)

    def test_ai_does_not_receive_raw_records_and_success_is_reused(self):
        report=self.report()
        content=json.dumps(ai.build_input(report))
        for field in ['source_record_id','raw_record_id','api_key','password','review_note']:self.assertNotIn(field,content)
        brief=self.brief(report)
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            again,reused=ai.enqueue(self.analyst,report)
        self.assertTrue(reused);self.assertEqual(again.pk,brief.pk);self.assertEqual(AIRequest.objects.count(),1)

    def test_same_input_new_report_keeps_its_own_brief_association(self):
        first=self.brief(self.report())
        second=self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            shared,reused=ai.enqueue(self.analyst,second)
        self.assertFalse(reused);self.assertNotEqual(shared.pk,first.pk)
        self.assertEqual(shared.report_id,second.pk)

    def test_same_report_changed_prompt_does_not_reuse_old_response(self):
        report=self.report();first=self.brief(report)
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            with patch.object(ai,'PROMPT_VERSION','unit-next'):
                changed,reused=ai.enqueue(self.analyst,report)
        self.assertFalse(reused);self.assertNotEqual(changed.pk,first.pk)

    def test_edit_preserves_original_and_rejects_stale_revision(self):
        brief=self.brief();original=brief.generated_text;raw=brief.raw_text
        ai.edit(self.analyst,brief.pk,'人工核查补充说明。',0)
        brief.refresh_from_db()
        self.assertEqual(brief.generated_text,original);self.assertEqual(brief.raw_text,raw)
        self.assertEqual(brief.revision,1);self.assertEqual(BriefRevision.objects.get(brief=brief).text,'人工核查补充说明。')
        with self.assertRaises(ValidationError):ai.edit(self.analyst,brief.pk,'旧页提交。',0)
        with self.assertRaises(DatabaseError),transaction.atomic():AIBrief.objects.filter(pk=brief.pk).update(raw_text='tampered')
        with self.assertRaises(DatabaseError),transaction.atomic():BriefRevision.objects.filter(brief=brief).update(text='tampered')

    def test_ai_timeout_retry_and_attempt_evidence(self):
        report=self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            brief,_=ai.enqueue(self.analyst,report);jobs.claim('brief')
            with patch('apps.reports.ai.call_model',side_effect=ai.AIError('timeout')):jobs.execute('brief',brief.pk)
            brief.refresh_from_db();self.assertEqual(brief.status,'failed');self.assertEqual(brief.error_code,'timeout');self.assertEqual(brief.raw_text,'')
            ai.retry(self.analyst,brief.pk);jobs.claim('brief')
            raw=valid_output(brief.input_summary)
            with patch('apps.reports.ai.call_model',return_value=(raw,ai.parse_output(raw,brief.input_summary),{'prompt_tokens':2,'completion_tokens':3},5)):jobs.execute('brief',brief.pk)
        brief.refresh_from_db();self.assertEqual(brief.status,'succeeded')
        self.assertEqual(list(brief.requests.values_list('status',flat=True)),['failed','succeeded'])

    def test_daily_budget_and_duplicate_pending_request_are_enforced(self):
        report=self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            brief,_=ai.enqueue(self.analyst,report)
            with self.assertRaises(ValidationError):ai.enqueue(self.analyst,report)
            with override_settings(AI_DAILY_CALL_LIMIT=0),self.assertRaises(ValidationError):ai.budget(self.other)

    def test_worker_recovery_marks_interrupted_jobs_failed(self):
        main,source=prepare_forms(self.analyst,{'kind':'statistics'})
        report=services.enqueue(self.analyst,main,source);jobs.claim('report')
        jobs.recover();report.refresh_from_db();self.assertEqual(report.error_code,'interrupted')

    def test_revoked_permissions_fail_before_any_model_request(self):
        report=self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            brief,_=ai.enqueue(self.analyst,report)
        self.analyst.groups.clear();jobs.claim('brief')
        with patch('apps.reports.ai.call_model') as call:jobs.execute('brief',brief.pk);call.assert_not_called()
        brief.refresh_from_db();self.assertEqual(brief.error_code,'permission')

    def test_post_routes_require_csrf_and_get_never_queues_work(self):
        report=self.report();before=Report.objects.count()
        secured=Client(enforce_csrf_checks=True);secured.force_login(self.analyst)
        for suffix in ['repeat','export','brief']:
            url=f'/reports/{report.pk}/{suffix}/'
            self.assertEqual(secured.post(url,{}).status_code,403)
            self.assertEqual(secured.get(url).status_code,405)
        self.assertEqual(Report.objects.count(),before)
        self.assertEqual(ExportJob.objects.count(),0);self.assertEqual(AIRequest.objects.count(),0)

    def test_retry_limit_and_changed_configuration_do_not_call_model(self):
        report=self.report()
        with patch('apps.reports.ai.load_config',return_value=CONFIG):
            brief,_=ai.enqueue(self.analyst,report);jobs.claim('brief')
            with patch('apps.reports.ai.call_model',side_effect=ai.AIError('timeout')):jobs.execute('brief',brief.pk)
            with override_settings(AI_MAX_ATTEMPTS=1),self.assertRaises(ValidationError):ai.retry(self.analyst,brief.pk)
            with patch.object(ai,'PROMPT_VERSION','changed'),self.assertRaises(ai.AIError):ai.retry(self.analyst,brief.pk)
        self.assertEqual(AIRequest.objects.count(),1)


class ExportPrecisionCleanupTests(SimpleTestCase):
    def test_large_identifiers_and_decimals_survive_excel_roundtrip(self):
        with tempfile.TemporaryDirectory() as root:
            target=Path(root)/'precision.xlsx'
            rows=[{'id':123456789012345678,'value':Decimal('1234567890.123456789'),'zero':0}]
            table=memory_table('exact','精度',[('id','标识'),('value','数值'),('zero','零')],rows)
            exports.write_xlsx(target,[table],{})
            wb=load_workbook(target,read_only=True)
            try:self.assertEqual(list(wb.worksheets[1].values)[1],('123456789012345678','1234567890.123456789',0))
            finally:wb.close()

    def test_failed_workbook_removes_only_its_own_internal_temp_files(self):
        from openpyxl.worksheet._writer import ALL_TEMP_FILES
        before=set(ALL_TEMP_FILES)
        with tempfile.TemporaryDirectory() as root:
            table=memory_table('bad','失败样本',[('value','值')],[{'value':'safe'},{'value':'x'*32768}])
            with self.assertRaises(exports.ExportError):exports.write_xlsx(Path(root)/'bad.xlsx',[table],{})
        self.assertEqual(set(ALL_TEMP_FILES),before)

    def test_multiple_valid_fact_references_are_accepted(self):
        summary={'facts':[{'id':f'F{i:03d}','text':f'分类 {i}：{i} 起'} for i in range(1,10)],'limitations':[{'id':'L001','text':'只展示排名摘录。'}]}
        output=json.loads(valid_output(summary));output['findings'][0]['refs']=[r['id'] for r in summary['facts']]
        rendered=ai.parse_output(json.dumps(output,ensure_ascii=False),summary)
        self.assertIn('分类 9：9 起',rendered)


@override_settings(ALLOWED_HOSTS=['testserver','127.0.0.1','localhost'],AI_CONFIG_FILE='')
class FlowReportTests(TestCase):
    @classmethod
    def setUpTestData(cls):flow_fixture.SpatialFlowTests.setUpTestData.__func__(cls)

    def test_flow_report_retains_zero_missing_and_exposure_ineligibility(self):
        day=date(2025,1,6)
        flow_fixture.SpatialFlowTests.observation(self,day,time(7),0)
        flow_fixture.SpatialFlowTests.observation(self,day,time(7,15),30)
        main,source=prepare_forms(self.analyst,{'kind':'flow','site':str(self.site.pk),'start':str(day),'end':str(day),'grain':'hour'})
        report=services.enqueue(self.analyst,main,source);jobs.claim('report');jobs.execute('report',report.pk)
        report.refresh_from_db();self.assertEqual(report.status,'succeeded')
        data=report.payload['analysis']
        self.assertEqual((data['coverage']['valid'],data['coverage']['observed_sum']),(2,30))
        self.assertIsNone(data['series'][7]['total']);self.assertFalse(report.payload['exposure']['eligible'])

    def test_observation_export_uses_half_open_clock_and_preserves_quality(self):
        day=date(2025,1,6)
        for clock,volume,quality in [(time(6,45),99,'valid'),(time(7),0,'valid'),(time(7,15),30,'excluded'),(time(10),99,'valid')]:
            flow_fixture.SpatialFlowTests.observation(self,day,clock,volume,quality=quality)
        form=exports.make_form('observations',{'site':str(self.site.pk),'date_from':str(day),'date_to':str(day),'raw_start_time':'07:00','raw_end_time':'10:00'})
        self.assertTrue(form.is_valid(),form.errors)
        values=list(exports.query_rows('observations',form).values_list('volume','quality_status'))
        self.assertEqual(values,[(0,'valid'),(30,'excluded')])
        bad=exports.make_form('observations',{'raw_start_time':'22:00','raw_end_time':'07:00'})
        self.assertFalse(bad.is_valid())
