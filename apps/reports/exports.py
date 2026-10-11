"""Bounded CSV/Excel writers over full query iterators, never UI page slices."""
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import io
import json
import zipfile
import csv
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import connection, transaction
from django.db.models import F, Func, FloatField
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from apps.analytics.query import consistent_read, data_version, plain
from apps.collisions.queries import QueryForm, filtered_entities
from apps.traffic.analysis import official_observations
from apps.core.audit import record_event
from apps.core.storage import storage_root
from .forms import ObservationExportForm
from .services import allow, lock_owner, source_metadata, describe
from .models import ExportJob
from .tables import Table, tables

CSV_DATA_ROWS=500000
XLSX_DATA_ROWS=1048575  # Excel row capacity minus the repeated column header.
KINDS={'crashes':'事故','vehicles':'车辆','persons':'人员','observations':'流量观测','report':'报告数据'}
ERRORS={'excel_cell_too_long':'文本超过 Excel 单元格容量，请改用 CSV 数据包导出，原始值未截断。',
        'invalid_text':'文本含无法写入电子表格的控制字符，请改用 CSV 数据包。','file_exists':'已有同名导出文件，请创建新的导出任务。',
        'interrupted':'上次导出中断，请重新提交。','permission':'账户权限已变化，导出未执行。',
        'invalid_scope':'筛选条件已失效，请重新设置。','internal':'导出未完成，请核对服务状态后重试。'}
FIELD_LABELS={'id':'内部编号','source_id':'来源目录编号','source_record_id':'来源记录编号','raw_record_id':'原始证据编号','raw_record__snapshot_id':'来源快照编号','source__name':'来源名称',
 'local_date':'纽约当地日期','local_time':'纽约当地时刻','occurred_at':'事故时刻（UTC）','observed_at':'观测时刻（UTC）','time_status':'时刻状态',
 'persons_injured':'事故级受伤人数','persons_killed':'事故级死亡人数','borough_id':'行政区编号','borough__chinese_name':'行政区',
 'on_street':'事发道路','off_street':'相邻道路','cross_street':'交叉街道／地址','collision_id':'事故内部编号','vehicle_id':'车辆内部编号',
 'origin':'记录来源类别','is_active':'有效状态','revision':'记录修订号','review_note':'核查备注','vehicle_type':'车辆类型','person_type':'人员类别',
 'injury':'伤害类别','age':'年龄','sex':'性别类别','volume':'15 分钟观测量（辆次）','quality_status':'质量状态','site_id':'采样地点编号',
 'site__street':'采样街道','site__direction':'采样方向','export_lon':'经度（WGS84）','export_lat':'纬度（WGS84）'}


class ExportError(Exception):
    pass


def make_form(kind,data):
    if kind in ('crashes','vehicles','persons'):return QueryForm(data,kind=kind)
    if kind=='observations':return ObservationExportForm(data)
    raise ValidationError('请选择有效的数据类型。')


def query_rows(kind,form):
    if kind!='observations':return filtered_entities(kind,form)
    qs=official_observations().order_by('local_date','local_time','pk')
    d=form.cleaned_data
    for key,lookup in [('site','site'),('borough','site__borough'),('date_from','local_date__gte'),('date_to','local_date__lte'),('raw_start_time','local_time__gte'),('raw_end_time','local_time__lt')]:
        if d.get(key) is not None:qs=qs.filter(**{lookup:d[key]})
    return qs


@transaction.atomic
def enqueue(actor,kind,fmt,form=None,report=None):
    allow(actor);lock_owner(actor,20261009)
    if fmt not in ('csv','xlsx'):raise ValidationError('请选择 CSV 或 Excel 格式。')
    if ExportJob.objects.filter(owner=actor,status__in=['queued','running']).exists():raise ValidationError('已有导出正在处理，请先查看任务。')
    if report:
        if report.owner_id!=actor.pk or report.status!='succeeded':raise ValidationError('请选择本人的已完成报告。')
        kind='report';conditions={}
    else:
        if kind not in KINDS or kind=='report' or form is None or not form.is_valid():raise ValidationError('请修正导出范围。')
        conditions=form.conditions()
    job=ExportJob.objects.create(owner=actor,kind=kind,format=fmt,report=report,conditions=conditions)
    record_event(actor=actor,action='export_requested',object_type='ExportJob',object_id=job.pk)
    return job


def path_for(job):
    root=(storage_root()/'var/exports').resolve()
    path=(root/(str(job.pk)+('.zip' if job.format=='csv' else '.xlsx'))).resolve()
    if path.parent!=root:raise ExportError('internal')
    return path


def text(value):
    if value is None:return ''
    if isinstance(value,bool):return '是' if value else '否'
    if isinstance(value,(date,datetime,time)):return value.isoformat()
    if isinstance(value,(list,dict)):return json.dumps(plain(value),ensure_ascii=False,sort_keys=True)
    return str(value)


def csv_value(value):
    result=text(value)
    if isinstance(value,str) and (result.lstrip().startswith(('=','+','-','@')) or result.startswith(('\t','\r','\n'))):
        return "'"+result
    return result


def excel_cell(sheet,value,header=False):
    if isinstance(value,(date,datetime,time,list,dict)):value=text(value)
    if isinstance(value,int) and not isinstance(value,bool) and abs(value)>=10**15:value=str(value)
    if isinstance(value,Decimal) and len(value.normalize().as_tuple().digits)>15:value=str(value)
    if isinstance(value,str):
        if len(value)>32767:raise ExportError('excel_cell_too_long')
        if any(ord(c)<32 and c not in '\t\r\n' for c in value):raise ExportError('invalid_text')
    cell=WriteOnlyCell(sheet,value=value)
    if isinstance(value,str):cell.data_type='s'
    if header:
        cell.font=Font(name='等线',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='175966')
    else:cell.font=Font(name='等线',size=10)
    cell.alignment=Alignment(vertical='top',wrap_text=header)
    return cell


def metadata_rows(metadata):
    names={'title':'内容','kind':'类型','created_at':'生成时间','conditions':'筛选条件','data_version':'数据版本','sources':'数据版本所涉来源快照','data_through':'来源数据截至','notes':'口径与质量说明','metrics':'报告确定指标','table_counts':'数据表行数','briefs':'简报版本信息','rule':'分析规则与参数','analysis_at':'原分析生成时间','template_version':'报告模板版本'}
    for key,value in metadata.items():
        if key=='sources':
            for i,source in enumerate(value,1):yield [f'来源 {i}',text(source)]
        elif key=='notes':
            for i,note in enumerate(value,1):yield [f'说明 {i}',note]
        else:yield [names.get(key,key),text(value)]


def write_csv(target,data_tables,metadata):
    counts={}
    with zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=4,allowZip64=True) as archive:
        for index,table in enumerate(data_tables,1):
            written=0;part=0;stream=None;wrapper=None
            try:
                for row in table.rows():
                    if written%CSV_DATA_ROWS==0:
                        if wrapper:wrapper.close()
                        part+=1
                        stream=archive.open(f'{index:02d}_{table.code}_{part:03d}.csv','w',force_zip64=True)
                        wrapper=io.TextIOWrapper(stream,encoding='utf-8-sig',newline='')
                        writer=csv.writer(wrapper);writer.writerow([x[1] for x in table.columns])
                    writer.writerow([csv_value(row.get(key)) for key,label in table.columns])
                    written+=1
                if not written:
                    stream=archive.open(f'{index:02d}_{table.code}_001.csv','w')
                    wrapper=io.TextIOWrapper(stream,encoding='utf-8-sig',newline='')
                    csv.writer(wrapper).writerow([x[1] for x in table.columns])
            finally:
                if wrapper:wrapper.close()
            if written!=table.count:raise ExportError('internal')
            counts[table.title]=written
        metadata['table_counts']=counts
        with archive.open('00_说明.csv','w') as stream:
            wrapper=io.TextIOWrapper(stream,encoding='utf-8-sig',newline='')
            writer=csv.writer(wrapper);writer.writerow(['项目','内容'])
            writer.writerows([[csv_value(v) for v in row] for row in metadata_rows(metadata)])
            wrapper.flush()
    return counts


def write_xlsx(target,data_tables,metadata):
    workbook=Workbook(write_only=True)
    try:
        return fill_xlsx(workbook,target,data_tables,metadata)
    finally:
        # openpyxl removes sheet XML after a successful save. On failure, close
        # and clean only writers created by this workbook, never scan tempdirs.
        for sheet in workbook.worksheets:
            writer=sheet._writer
            if writer is not None and Path(writer.out).is_file():
                try:
                    if not sheet.closed:sheet.close()
                finally:
                    writer.cleanup()
        workbook.close()


def fill_xlsx(workbook,target,data_tables,metadata):
    counts={}
    info=workbook.create_sheet('说明')
    info.column_dimensions['A'].width=28;info.column_dimensions['B'].width=100
    info.append([excel_cell(info,v,True) for v in ['项目','内容']])
    for index,table in enumerate(data_tables,1):
        written=0;part=0;sheet=None
        def new_sheet():
            nonlocal part
            part+=1
            s=workbook.create_sheet(f'{index:02d}_{table.code[:19]}_{part:02d}')
            s.freeze_panes='A2'
            s.print_title_rows='1:1'
            s.sheet_properties.pageSetUpPr.fitToPage=True
            s.page_setup.orientation='landscape';s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
            for j,col in enumerate(table.columns,1):s.column_dimensions[get_column_letter(j)].width=22 if len(table.columns)<12 else 18
            s.append([excel_cell(s,label,True) for key,label in table.columns])
            return s
        for row in table.rows():
            if written%XLSX_DATA_ROWS==0:
                if sheet:sheet.auto_filter.ref=f'A1:{get_column_letter(len(table.columns))}{XLSX_DATA_ROWS+1}'
                sheet=new_sheet()
            sheet.append([excel_cell(sheet,row.get(key)) for key,label in table.columns])
            written+=1
        if sheet is None:sheet=new_sheet()
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(table.columns))}{(written-1)%XLSX_DATA_ROWS+2 if written else 1}'
        if written!=table.count:raise ExportError('internal')
        counts[table.title]=written
    metadata['table_counts']=counts
    for row in metadata_rows(metadata):
        cells=[excel_cell(info,v) for v in row]
        for c in cells:c.alignment=Alignment(wrap_text=True,vertical='top')
        info.append(cells)
    workbook.save(target)
    return counts


def source_table(job):
    form=make_form(job.kind,job.conditions)
    if not form.is_valid():raise ExportError('invalid_scope')
    qs=query_rows(job.kind,form)
    fields=[f.attname for f in qs.model._meta.local_fields if f.name!='geom']
    fields+=['source__name','raw_record__snapshot_id']
    if job.kind=='crashes':
        fields+=['borough__chinese_name','export_lon','export_lat']
        qs=qs.annotate(export_lon=Func(F('geom'),function='ST_X',output_field=FloatField()),export_lat=Func(F('geom'),function='ST_Y',output_field=FloatField()))
    if job.kind=='observations':fields+=['site__street','site__direction']
    columns=[(key,FIELD_LABELS.get(key,key)) for key in fields]
    return Table(job.kind,KINDS[job.kind]+'完整筛选记录',columns,qs.count(),lambda:qs.values(*fields).iterator(chunk_size=4000))


def compute(pk):
    target=None;temporary=None
    temporary_owned=False;target_owned=False
    try:
        with consistent_read():
            with connection.cursor() as cursor:
                cursor.execute('SELECT pg_try_advisory_xact_lock_shared(%s,%s)',[20261004,2])
                if not cursor.fetchone()[0]:return False
            job=ExportJob.objects.select_related('owner','report__risk_run','report__forecast_run').get(pk=pk)
            allow(job.owner)
            if job.status!='running':return True
            target=path_for(job);target.parent.mkdir(parents=True,exist_ok=True)
            temporary=target.with_suffix(target.suffix+'.part')
            if target.exists() or temporary.exists():raise ExportError('file_exists')
            temporary.touch(exist_ok=False);temporary_owned=True
            if job.kind=='report':
                report=job.report
                data_tables=list(tables(report))
                metadata={'title':report.title,'kind':report.get_kind_display(),'created_at':report.finished_at,'conditions':report.payload['scope'],
                          'data_version':report.data_version,'sources':report.payload['sources'],'data_through':report.payload['data_through'],'notes':list(report.payload['notes']),'metrics':report.payload['metrics']}
                metadata.update(template_version=report.template_version,rule=report.payload.get('rule'),analysis_at=report.payload.get('analysis_at'))
                briefs=list(report.briefs.filter(status='succeeded').order_by('created_at'))
                if briefs:
                    from .tables import memory_table
                    metadata['briefs']=[dict(model=b.model,prompt_version=b.prompt_version,generated_at=b.finished_at,edited_at=b.edited_at,revision=b.revision) for b in briefs]
                    rows=[dict(model=b.model,generated_at=b.finished_at,revision=b.revision,original=b.generated_text,current=b.edited_text if b.revision else b.generated_text) for b in briefs]
                    data_tables.append(memory_table('briefs','报告简报文本',[('model','模型'),('generated_at','生成时间'),('revision','编辑版本'),('original','模型生成文本'),('current','当前文本')],rows))
            else:
                version=data_version();sources=source_metadata(version)
                data_tables=[source_table(job)]
                domain={'crashes':'crashes','vehicles':'vehicles','persons':'persons','observations':'traffic'}[job.kind]
                metadata={'title':KINDS[job.kind]+'筛选数据','kind':KINDS[job.kind],'created_at':timezone.now(),'conditions':job.conditions,
                          'data_version':version,'sources':sources,'data_through':max((s['data_through'] for s in sources if s['kind']==domain and s['data_through']),default=None),
                          'notes':['导出包含所选筛选条件的全部记录，不受网页分页影响。','缺失保留空单元格，0 是真实零；日期和时刻采用字段注明的纽约当地时间或 UTC。','字段采用业务列名，定义见系统数据字典。流量 conflict/excluded 等行可作为证据导出，不因此取得统计资格。']}
            metadata['notes']+=['CSV 文本的公式前导字符加单引号；Excel 文本作为字符串写入，原始数据库值不变。超过十五位有效数字的整数和小数在 Excel 中按文本保留。','CSV 每五十万行拆文件；Excel 达工作表行容量时拆表，每份重复表头。']
            counts=write_csv(temporary,data_tables,metadata) if job.format=='csv' else write_xlsx(temporary,data_tables,metadata)
            temporary.replace(target)
            target_owned=True
            job.row_count=sum(counts.values());job.file_size=target.stat().st_size;job.relative_path=target.relative_to(storage_root()).as_posix()
            job.metadata=plain(metadata);job.status='succeeded';job.finished_at=timezone.now();job.error_code='';job.save()
            record_event(actor=job.owner,action='export_created',object_type='ExportJob',object_id=pk,details={'count':job.row_count})
        return True
    except Exception:
        # Exact task-owned files only; never walk or remove a directory.
        root=(storage_root()/'var/exports').resolve()
        for path,owned in ((temporary,temporary_owned),(target,target_owned)):
            if owned and path and path.resolve().parent==root and path.name.startswith(str(pk)) and path.exists():path.unlink()
        raise
